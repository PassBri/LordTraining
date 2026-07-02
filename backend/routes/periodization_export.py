"""Exportar la Tabla Compacta de Periodización a un archivo Excel (.xlsx) real."""
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import List, Optional
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

router = APIRouter(prefix="/periodization", tags=["periodization"])


def get_current_user_dep():
    from server import get_current_user
    return get_current_user


class ExcelRow(BaseModel):
    semana: Optional[str] = ""
    macro: Optional[str] = ""
    fase: Optional[str] = ""
    meso: Optional[str] = ""
    tipo: Optional[str] = ""
    volumen: Optional[float] = None
    intensidad: Optional[float] = None
    carga: Optional[float] = None
    sesiones: Optional[int] = None
    acwr: Optional[float] = None
    estado: Optional[str] = ""


class ExcelReq(BaseModel):
    plan_name: Optional[str] = "Plan de Periodización"
    rows: List[ExcelRow] = []


_HEADERS = [
    ("semana", "Semana"), ("macro", "Macrociclo"), ("fase", "Fase"),
    ("meso", "Mesociclo"), ("tipo", "Tipo Micro"), ("volumen", "Volumen %"),
    ("intensidad", "Intensidad %"), ("carga", "Carga %"), ("sesiones", "Sesiones"),
    ("acwr", "ACWR"), ("estado", "Estado de carga"),
]

_STATUS_FILL = {
    "optimal": "D1FAE5", "caution": "FEF3C7", "overload": "FEE2E2", "detraining": "FEE2E2",
}


@router.post("/export-excel")
async def export_excel(payload: ExcelReq, user=Depends(get_current_user_dep())):
    wb = Workbook()
    ws = wb.active
    ws.title = "Periodización"

    navy = "0A0F1E"
    lime = "CCFF00"
    thin = Side(style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Título
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_HEADERS))
    tcell = ws.cell(row=1, column=1, value=payload.plan_name)
    tcell.font = Font(bold=True, size=14, color="FFFFFF")
    tcell.fill = PatternFill("solid", fgColor=navy)
    tcell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    # Encabezados
    hr = 2
    for col, (_, label) in enumerate(_HEADERS, start=1):
        c = ws.cell(row=hr, column=col, value=label)
        c.font = Font(bold=True, size=10, color=navy)
        c.fill = PatternFill("solid", fgColor=lime)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    # Datos
    for i, row in enumerate(payload.rows, start=hr + 1):
        vals = row.dict()
        for col, (key, _) in enumerate(_HEADERS, start=1):
            v = vals.get(key)
            c = ws.cell(row=i, column=col, value=v if v is not None else "")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
            c.font = Font(size=10)
        # Colorear la fila 'estado'
        fill_hex = _STATUS_FILL.get((row.estado or "").lower())
        if fill_hex:
            ws.cell(row=i, column=len(_HEADERS)).fill = PatternFill("solid", fgColor=fill_hex)

    widths = [10, 22, 20, 20, 18, 12, 14, 10, 10, 10, 16]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=hr, column=col).column_letter].width = w
    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = (payload.plan_name or "plan").replace(" ", "_")[:40]
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{safe}.xlsx"'},
    )
